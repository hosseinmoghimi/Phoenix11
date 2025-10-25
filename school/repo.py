from .models import School,Course,CourseClass,Teacher,Student,Major,Session,StudentInSession
from .apps import APP_NAME
from .enums import *
from log.repo import LogRepo 
from django.db.models import Q
from django.shortcuts import reverse
from authentication.repo import PersonRepo
from accounting.repo import InvoiceLineItemUnitRepo
from utility.num import filter_number
from utility.calendar import PersianCalendar
from utility.constants import FAILED,SUCCEED
from utility.log import leolog
from .enums import *


class SchoolRepo():
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=School.objects.filter(id=0)
        profile=PersonRepo(request=request).me
        if profile is not None:
            if request.user.has_perm(APP_NAME+".view_account"):
                self.objects=School.objects
                self.my_accounts=self.objects 
    def list(self,*args, **kwargs):
        objects=self.objects
        if "search_for" in kwargs:
            search_for=kwargs["search_for"]
            objects=objects.filter(Q(name__contains=search_for) | Q(code=search_for)  )
        if "parent_id" in kwargs:
            parent_id=kwargs["parent_id"]
            objects=objects.filter(parent_id=parent_id)  
        return objects.all()
        
    def school(self,*args, **kwargs):
        if "school_id" in kwargs and kwargs["school_id"] is not None:
            return self.objects.filter(pk=kwargs['school_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
        
    def add_school(self,*args,**kwargs):
        result,message,school=FAILED,"",None
        if len(School.objects.filter(name=kwargs["name"]))>0:
            message='نام تکراری برای آموزشگاه جدید'
            return FAILED,message,None
        if len(School.objects.filter(person_account_id=kwargs["person_account_id"]))>0:
            message='حساب تکراری برای آموزشگاه جدید'
            return FAILED,message,None
        if not self.request.user.has_perm(APP_NAME+".add_school"):
            message="دسترسی غیر مجاز"
            return result,message,school

        school=School() 
        if 'person_account_id' in kwargs:
            school.person_account_id=kwargs["person_account_id"]
        if 'name' in kwargs:
            school.name=kwargs["name"]
          
        (result,message,school)=school.save()
        return result,message,school
 

class StudentRepo():
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=Student.objects.filter(id=0)
        profile=PersonRepo(request=request).me
        if profile is not None:
            if request.user.has_perm(APP_NAME+".view_account"):
                self.objects=Student.objects
                self.my_accounts=self.objects 
    def import_students(self,*args, **kwargs):
        result,message,students=FAILED,'',[]
        excel_file=kwargs['excel_file']
        import openpyxl 

        wb = openpyxl.load_workbook(excel_file)
        try:
            ws = wb['students']
        except:
            message='فایل شما برگه دانش آموزان ندارد.'
            return result,message,None
        count=kwargs['count']
        try:
            count=int(ws.cell(row=1, column=2).value)
        except:
            message='فایل برگه دانش آموزان ، تعداد ندارد.'
            return result,message,None 
        from .constants import EXCEL_MAJORS_DATA_START_ROW,EXCEL_STUDENTS_DATA_START_ROW,EXCEL_TEACHERS_DATA_START_ROW

        students_to_import=[]
        START_ROW=EXCEL_STUDENTS_DATA_START_ROW

        modified=added=0 
        for i in range(START_ROW,count+START_ROW):
            i=str(i) 
            # student['id']=ws['A'+str(i)].value
            iiiddd=ws['B'+i].value
            if iiiddd is not None:
                id=int(ws['B'+i].value)
                prefix=str(ws['C'+i].value)
                last_name=str(ws['D'+i].value)
                first_name=str(ws['E'+i].value)
                father_name=str(ws['F'+i].value)
                melli_code=str(ws['G'+i].value)
                birth_date=str(ws['H'+i].value)
                birth_location=str(ws['I'+i].value)
                leolog(prefix=prefix,first_name=first_name,last_name=last_name,melli_code=melli_code,father_name=father_name)
                from authentication.repo import Person,PersonRepo
                from accounting.repo import PersonAccount
                person=PersonRepo(request=self.request).person(melli_code=melli_code)
                if person is None:
                    person=Person()
                else:
                    modified=modified+1
                if prefix is not None and not prefix=='':
                    person.prefix=prefix
                if first_name is not None and not first_name=='':
                    person.first_name=first_name
                if last_name is not None and not last_name=='':
                    person.last_name=last_name
                if melli_code is not None and not melli_code=='':
                    person.melli_code=melli_code
                if birth_date is not None and not birth_date=='':
                    person.birth_date=birth_date
                if birth_location is not None and not birth_location=='':
                    person.birth_location=birth_location
                if father_name is not None and not father_name=='':
                    person.father_name=father_name
                result,message,person=person.save()
                if result==FAILED:
                    return result,message,[]
                person_account=PersonAccount()
                person_account.person=person
                person_category_id=kwargs['person_category_id']
                person_account.person_category_id=person_category_id
                person_account.save()
                student=Student.objects.filter(person_account__person_id=person.id).first()
                if student is None:
                    student=Student(person_account_id=person_account.id)
                    student.save()
                    added=added+1
                
        result=SUCCEED
        message=f"""{added} دانش آموز اضافه شد.
                    <br>
                    {modified} دانش آموز ویرایش شد. """
        students=self.list()
        result=SUCCEED
        message+='<br>'+'با موفقیت بازیابی شد.  '
        return result,message,students
    def list(self,*args, **kwargs):
        objects=self.objects
        if "search_for" in kwargs:
            search_for=kwargs["search_for"]
            objects=objects.filter(Q(name__contains=search_for) | Q(code=search_for)  )
        if "parent_id" in kwargs:
            parent_id=kwargs["parent_id"]
            objects=objects.filter(parent_id=parent_id)  
        return objects.all()
        
    def student(self,*args, **kwargs):
        if "student_id" in kwargs and kwargs["student_id"] is not None:
            return self.objects.filter(pk=kwargs['student_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
        
    def add_student(self,*args,**kwargs):
        result,message,student=FAILED,"",None 
        if len(Student.objects.filter(person_account_id=kwargs["person_account_id"]))>0:
            message='حساب تکراری برای دانش آموز جدید'
            return FAILED,message,None
        if not self.request.user.has_perm(APP_NAME+".add_student"):
            message="دسترسی غیر مجاز"
            return result,message,student

        student=Student() 
        if 'person_account_id' in kwargs:
            student.person_account_id=kwargs["person_account_id"]
        if 'name' in kwargs:
            student.name=kwargs["name"]
          
        (result,message,student)=student.save()
        return result,message,student

 
class TeacherRepo():
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=Teacher.objects.filter(id=0)
        profile=PersonRepo(request=request).me
        if profile is not None:
            if request.user.has_perm(APP_NAME+".view_account"):
                self.objects=Teacher.objects
                self.my_accounts=self.objects 
    def list(self,*args, **kwargs):
        objects=self.objects
        if "search_for" in kwargs:
            search_for=kwargs["search_for"]
            objects=objects.filter(Q(name__contains=search_for) | Q(code=search_for)  )
        if "parent_id" in kwargs:
            parent_id=kwargs["parent_id"]
            objects=objects.filter(parent_id=parent_id)  
        return objects.all()
        
    def teacher(self,*args, **kwargs):
        if "teacher_id" in kwargs and kwargs["teacher_id"] is not None:
            return self.objects.filter(pk=kwargs['teacher_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
        
    def add_teacher(self,*args,**kwargs):
        result,message,teacher=FAILED,"",None 
        if len(Teacher.objects.filter(person_account_id=kwargs["person_account_id"]))>0:
            message='حساب تکراری برای دبیر جدید'
            return FAILED,message,None
        if not self.request.user.has_perm(APP_NAME+".add_teacher"):
            message="دسترسی غیر مجاز"
            return result,message,teacher

        teacher=Teacher() 
        if 'person_account_id' in kwargs:
            teacher.person_account_id=kwargs["person_account_id"]
        if 'name' in kwargs:
            teacher.name=kwargs["name"]
          
        (result,message,teacher)=teacher.save()
        return result,message,teacher

 
class MajorRepo():
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=Major.objects.filter(id=0)
        profile=PersonRepo(request=request).me
        if profile is not None:
            if request.user.has_perm(APP_NAME+".view_account"):
                self.objects=Major.objects
                self.my_accounts=self.objects 
    def list(self,*args, **kwargs):
        objects=self.objects
        if "search_for" in kwargs:
            search_for=kwargs["search_for"]
            objects=objects.filter(Q(name__contains=search_for) | Q(code=search_for)  )
        if "parent_id" in kwargs:
            parent_id=kwargs["parent_id"]
            objects=objects.filter(parent_id=parent_id)  
        return objects.all()
        
    def major(self,*args, **kwargs):
        if "major_id" in kwargs and kwargs["major_id"] is not None:
            return self.objects.filter(pk=kwargs['major_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
        
    def add_major(self,*args,**kwargs):
        result,message,major=FAILED,"",None
        if len(Major.objects.filter(title=kwargs["title"]))>0:
            message='نام تکراری برای رشته جدید'
            return FAILED,message,None
      
        if not self.request.user.has_perm(APP_NAME+".add_major"):
            message="دسترسی غیر مجاز"
            return result,message,major

        major=Major() 
        if 'title' in kwargs:
            major.title=kwargs["title"]
        
          
        (result,message,major)=major.save()
        return result,message,major

 
class CourseRepo():
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=Course.objects.filter(id=0)
        profile=PersonRepo(request=request).me
        if profile is not None:
            if request.user.has_perm(APP_NAME+".view_account"):
                self.objects=Course.objects
                self.my_accounts=self.objects 
    def list(self,*args, **kwargs):
        objects=self.objects
        if "search_for" in kwargs:
            search_for=kwargs["search_for"]
            objects=objects.filter(Q(name__contains=search_for) | Q(code=search_for)  )
        if "parent_id" in kwargs:
            parent_id=kwargs["parent_id"]
            objects=objects.filter(parent_id=parent_id)  
        return objects.all()
        
    def course(self,*args, **kwargs):
        if "course_id" in kwargs and kwargs["course_id"] is not None:
            return self.objects.filter(pk=kwargs['course_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
        
    def add_course(self,*args,**kwargs):
        result,message,course=FAILED,"",None
        if not self.request.user.has_perm(APP_NAME+".add_course"):
            message="دسترسی غیر مجاز"
            return result,message,course

        course=Course()
        if 'title' in kwargs:
            course.title=kwargs["title"]
        if 'parent_id' in kwargs:
            if kwargs["parent_id"]>0:
                course.parent_id=kwargs["parent_id"]
        if 'color' in kwargs:
            course.color=kwargs["color"]
        if 'code' in kwargs:
            course.code=kwargs["code"]
        if 'priority' in kwargs:
            course.priority=kwargs["priority"]
        if 'type' in kwargs:
            course.type=kwargs["type"]

            
        if 'parent_code' in kwargs:
            parent_code= kwargs["parent_code"]
            parent=Account.objects.filter(code=parent_code).first()
            if parent is not None:
                course.parent_id=parent.id

        if 'nature' in kwargs:
            course.nature=kwargs["nature"]
        (result,message,course)=course.save()
        return result,message,course
 

class CourseClassRepo():
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=CourseClass.objects.filter(id=0)
        me_person=PersonRepo(request=request).me
        if me_person is not None:
            if request.user.has_perm(APP_NAME+".view_courseclass"):
                self.objects=CourseClass.objects
                self.my_accounts=self.objects 
            else:
                me_student=StudentRepo(request=request).me
                if me_student is not None:
                    self.objects=me_student.courseclass_set.all()
                
                else:
                    me_teacher=TeacherRepo(request=request).me
                    if me_teacher is not None:
                        self.objects=me_teacher.courseclass_set.all()
    def list(self,*args, **kwargs):
        objects=self.objects
        if "search_for" in kwargs:
            search_for=kwargs["search_for"]
            objects=objects.filter(Q(name__contains=search_for) | Q(code=search_for)  )
        if "parent_id" in kwargs:
            parent_id=kwargs["parent_id"]
            objects=objects.filter(parent_id=parent_id)  
        if "major_id" in kwargs:
            major_id=kwargs["major_id"]
            courses=CourseRepo(request=self.request).list(major_id=major_id)
            courses_ids=[]
            for course in courses:
                courses_ids.append(course.id)
            objects=objects.filter(course_id__in=courses_ids) 
        if "school_id" in kwargs:
            objects=objects.filter(school_id=kwargs['school_id'])  
        if "course_id" in kwargs:
            objects=objects.filter(course_id=kwargs['course_id'])  
        return objects.all()
        
    def course_class(self,*args, **kwargs):
        if "course_class_id" in kwargs and kwargs["course_class_id"] is not None:
            return self.objects.filter(pk=kwargs['course_class_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
        
    def add_course_class(self,*args,**kwargs):
        result,message,course_class=FAILED,"",None
        if not self.request.user.has_perm(APP_NAME+".add_course_class"):
            message="دسترسی غیر مجاز"
            return result,message,course_class

        course_class=CourseClass()
        if 'name' in kwargs:
            course_class.name=kwargs["name"]
        if 'parent_id' in kwargs:
            if kwargs["parent_id"]>0:
                course_class.parent_id=kwargs["parent_id"]
        if 'color' in kwargs:
            course_class.color=kwargs["color"]
        if 'code' in kwargs:
            course_class.code=kwargs["code"]
        if 'priority' in kwargs:
            course_class.priority=kwargs["priority"]
        if 'type' in kwargs:
            course_class.type=kwargs["type"]

            
        if 'parent_code' in kwargs:
            parent_code= kwargs["parent_code"]
            parent=Account.objects.filter(code=parent_code).first()
            if parent is not None:
                course_class.parent_id=parent.id

        if 'nature' in kwargs:
            course_class.nature=kwargs["nature"]
        (result,message,course_class)=course_class.save()
        return result,message,course_class


class SessionRepo():
    
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=Session.objects.filter(id=0)
        me_person=PersonRepo(request=request).me
        if me_person is not None:
            if request.user.has_perm(APP_NAME+".view_courseclass"):
                self.objects=Session.objects
                self.my_accounts=self.objects 
            else:
                me_student=StudentRepo(request=request).me
                if me_student is not None:
                    self.objects=me_student.courseclass_set.all()
                
                else:
                    me_teacher=TeacherRepo(request=request).me
                    if me_teacher is not None:
                        self.objects=me_teacher.courseclass_set.all()
    def list(self,*args, **kwargs):
        objects=self.objects
        if "search_for" in kwargs:
            search_for=kwargs["search_for"]
            objects=objects.filter(Q(name__contains=search_for) | Q(code=search_for)  )
        if "course_class_id" in kwargs:
            course_class_id=kwargs["course_class_id"]
            objects=objects.filter(course_class_id=course_class_id)  
           
        if "course_id" in kwargs:
            objects=objects.filter(course_class__course_id=kwargs['course_id'])  
        return objects.all()
        
    def session(self,*args, **kwargs):
        if "session_id" in kwargs and kwargs["session_id"] is not None:
            return self.objects.filter(pk=kwargs['session_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
        
    def add_session(self,*args,**kwargs):
        result,message,session=FAILED,"",None
        if not self.request.user.has_perm(APP_NAME+".add_session"):
            message="دسترسی غیر مجاز"
            return result,message,session
        from django.utils import timezone
        now=timezone.now()
        session=Session()
        if 'course_class_id' in kwargs:
            session.course_class_id=kwargs["course_class_id"] 
        if 'session_no' in kwargs:
            session.session_no=kwargs["session_no"]
        if 'start_datetime' in kwargs:
            session.start_datetime=kwargs["start_datetime"]
        else:
             session.start_datetime=now
        if 'end_datetime' in kwargs:
            session.end_datetime=kwargs["end_datetime"]
        else:
             session.end_datetime=now 
              
        (result,message,session)=session.save()
        student_in_session_repo=StudentInSessionRepo(request=self.request)
        for student in session.course_class.students.all():
            student_in_session_repo.add_student_in_session(student_id=student.id,session_id=session.id)
        return result,message,session


class StudentInSessionRepo():
    
    def __init__(self,request,*args, **kwargs):
        self.me=None
        self.my_accounts=[]
        self.request=request
        self.objects=StudentInSession.objects.filter(id=0)
        me_person=PersonRepo(request=request).me
        if me_person is not None:
            if request.user.has_perm(APP_NAME+".view_courseclass"):
                self.objects=StudentInSession.objects
                self.my_accounts=self.objects 
            else:
                me_student=StudentRepo(request=request).me
                if me_student is not None:
                    self.objects=me_student.courseclass_set.all()
                
                else:
                    me_teacher=TeacherRepo(request=request).me
                    if me_teacher is not None:
                        self.objects=me_teacher.courseclass_set.all()
    
    def list(self,*args, **kwargs):
        objects=self.objects 
        if "session_id" in kwargs:
            session_id=kwargs["session_id"]
            objects=objects.filter(session_id=session_id) 
        if "student_id" in kwargs:
            student_id=kwargs["student_id"]
            objects=objects.filter(student_id=student_id)  
          
        return objects.all()
        
    def student_in_session(self,*args, **kwargs):
        if "student_in_session_id" in kwargs and kwargs["student_in_session_id"] is not None:
            return self.objects.filter(pk=kwargs['student_in_session_id']).first()  
        if "pk" in kwargs and kwargs["pk"] is not None:
            return self.objects.filter(pk=kwargs['pk']).first() 
        if "id" in kwargs and kwargs["id"] is not None:
            return self.objects.filter(pk=kwargs['id']).first() 
        
    def add_student_in_session(self,*args,**kwargs):
        result,message,student_in_session=FAILED,"",None
        if not self.request.user.has_perm(APP_NAME+".add_student_in_session"):
            message="دسترسی غیر مجاز"
            return result,message,student_in_session
        
        StudentInSession.objects.filter(student_id=kwargs["student_id"]).filter(session_id=kwargs["session_id"]).delete()

        student_in_session=StudentInSession()
       
        
        if 'student_id' in kwargs:
            student_in_session.student_id=kwargs["student_id"]
        if 'session_id' in kwargs:
            student_in_session.session_id=kwargs["session_id"]

        if 'status' in kwargs:
            student_in_session.status=kwargs["status"]
        if 'score' in kwargs:
            student_in_session.score=kwargs["score"]
        if 'description' in kwargs:
            student_in_session.description=kwargs["description"]

        student_in_session.save()
        if student_in_session.id is not None:    
            (result,message,student_in_session)=(SUCCEED,'با موفقیت ذخیره شد.',student_in_session)
        return result,message,student_in_session
